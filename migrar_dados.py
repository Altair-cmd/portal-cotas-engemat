"""
migrar_dados.py — Importa os dados da planilha Excel para o Supabase
Execute UMA VEZ para popular o banco com os dados atuais.

Uso:
  python migrar_dados.py "caminho\\para\\CONTROLE E ACOMPANHAMENTO JOVENS E PCD.xlsx"
"""
import sys
import math
import datetime
import openpyxl
from supabase import create_client

# ─── Configuração ───────────────────────────────────────────────
SUPABASE_URL = ""
SUPABASE_KEY = ""

# ─── Helpers ────────────────────────────────────────────────────
def _num(v):
    try:
        f = float(v) if v is not None else 0.0
        return 0.0 if math.isnan(f) else f
    except (ValueError, TypeError):
        return 0.0


def ler_planilha(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    if "RESUMO" not in wb.sheetnames:
        raise ValueError("Aba 'RESUMO' não encontrada na planilha.")
    ws = wb["RESUMO"]
    obras = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        cc = row[2]
        if not cc:
            continue
        nec_apr   = int(_num(row[7]))
        atual_apr = int(_num(row[8]))
        nec_pcd   = int(_num(row[11]))
        atual_pcd = int(_num(row[12]))
        def_apr   = atual_apr - nec_apr
        def_pcd   = atual_pcd - nec_pcd

        def status(d, nec):
            if nec == 0:   return "Dentro da cota"
            if d < 0:      return "Abaixo da cota"
            if d > 0:      return "Acima da cota"
            return "Dentro da cota"

        obras.append({
            "tipo":          str(row[0] or "").strip(),
            "cnpj":          str(row[1] or "").strip(),
            "obra":          str(cc).strip(),
            "emp_sienge":    str(row[3] or "").strip(),
            "emp_dom":       str(row[4] or "").strip(),
            "cc_dom":        str(row[5] or "").strip(),
            "qtd_func":      int(_num(row[6])),
            "nec_apr":       nec_apr,
            "atual_apr":     atual_apr,
            "def_apr":       def_apr,
            "status_apr":    status(def_apr, nec_apr),
            "nec_pcd":       nec_pcd,
            "atual_pcd":     atual_pcd,
            "def_pcd":       def_pcd,
            "status_pcd":    status(def_pcd, nec_pcd),
            "atualizado_em": datetime.datetime.now().isoformat(),
        })
    return obras


# ─── Main ───────────────────────────────────────────────────────
def main():
    if len(sys.argv) < 2:
        caminho = input("Caminho da planilha .xlsx: ").strip().strip('"')
    else:
        caminho = sys.argv[1]

    print(f"\n📂 Lendo planilha: {caminho}")
    obras = ler_planilha(caminho)
    print(f"✅ {len(obras)} obras encontradas na aba RESUMO")

    print("\n🔌 Conectando ao Supabase...")
    sb = create_client(SUPABASE_URL, SUPABASE_KEY)

    print("🗑️  Limpando dados anteriores...")
    sb.table("obras").delete().neq("id", "00000000-0000-0000-0000-000000000000").execute()

    print("📤 Inserindo obras no banco...")
    BATCH = 50
    for i in range(0, len(obras), BATCH):
        batch = obras[i:i + BATCH]
        sb.table("obras").insert(batch).execute()
        print(f"   {min(i + BATCH, len(obras))}/{len(obras)} obras inseridas...")

    print(f"\n🎉 Migração concluída! {len(obras)} obras salvas no Supabase.")
    print("   Agora os dados estão disponíveis no portal online.")


if __name__ == "__main__":
    main()
